import csv
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, DecimalField, F, Max, Q, Sum
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from accounts.permissions import manager_required
from catalog.models import Product
from customers.models import Customer
from inventory.models import StockLevel, StockTransfer
from sales.models import Payment, Sale, SaleItem

from .models import Branch, Notification

MONEY = DecimalField(max_digits=12, decimal_places=2)


@login_required
def home(request):
    profile = getattr(request.user, "profile", None)
    branch = profile.branch if profile else None
    is_manager = bool(profile and profile.is_manager)

    branch_id = request.GET.get("branch")
    if is_manager and branch_id:
        branch = Branch.objects.filter(id=branch_id).first() or branch

    hour = timezone.localtime().hour
    if hour < 12:
        greeting = "Good morning"
    elif hour < 18:
        greeting = "Good afternoon"
    else:
        greeting = "Good evening"

    context = {"active_nav": "home", "branch": branch, "greeting": greeting}
    if is_manager:
        context["branches"] = Branch.objects.filter(is_active=True, is_warehouse=False)

    if branch:
        today = timezone.localdate()
        today_agg = Sale.objects.filter(
            branch=branch, status=Sale.STATUS_COMPLETED, completed_at__date=today
        ).aggregate(count=Count("id"), total=Sum("total"))

        stock_qs = StockLevel.objects.filter(branch=branch)
        low_stock_products = (
            stock_qs.filter(quantity__lte=F("reorder_threshold"))
            .select_related("product")
            .order_by("quantity")[:5]
        )
        out_of_stock_count = stock_qs.filter(quantity=0).count()

        open_sale = (
            Sale.objects.filter(cashier=request.user, status=Sale.STATUS_OPEN)
            .prefetch_related("items")
            .first()
        )

        recent_sales = (
            Sale.objects.filter(branch=branch, status=Sale.STATUS_COMPLETED)
            .select_related("customer", "cashier")
            .order_by("-completed_at")[:5]
        )

        # Top sellers, last 30 days at this branch — a quick glance at what's
        # moving without having to open the full Reports page.
        top_sellers_since = today - timedelta(days=30)
        top_sellers = list(
            SaleItem.objects.filter(
                sale__branch=branch, sale__status=Sale.STATUS_COMPLETED,
                sale__completed_at__date__gte=top_sellers_since, product__isnull=False,
            )
            .values("product__id", "product__sku", "product__name")
            .annotate(
                units=Coalesce(Sum("quantity"), Decimal("0"), output_field=MONEY),
                revenue=Coalesce(Sum("line_total"), Decimal("0"), output_field=MONEY),
            )
            .order_by("-revenue")[:5]
        )

        context.update({
            "pending_transfers": StockTransfer.objects.filter(
                from_branch=branch, status=StockTransfer.STATUS_PENDING
            ).select_related("product", "to_branch").count(),
            "today_sale_count": today_agg["count"] or 0,
            "today_sale_total": today_agg["total"] or 0,
            "low_stock_products": low_stock_products,
            "low_stock_total": stock_qs.filter(quantity__lte=F("reorder_threshold")).count(),
            "out_of_stock_count": out_of_stock_count,
            "open_sale": open_sale,
            "recent_sales": recent_sales,
            "top_sellers": top_sellers,
            "customer_count": Customer.objects.count(),
        })

    return render(request, "core/dashboard.html", context)


@login_required
def global_search(request):
    q = (request.GET.get("q") or "").strip()
    products, customers = [], []
    if q:
        product_qs = Product.objects.filter(is_active=True).filter(
            Q(name__icontains=q) | Q(sku__icontains=q)
        ).order_by("name")[:5]
        products = [
            {"id": p.id, "sku": p.sku, "name": p.name, "url": f"{reverse('inventory:list')}?q={p.sku}"}
            for p in product_qs
        ]
        customer_qs = Customer.objects.filter(
            Q(name__icontains=q) | Q(phone__icontains=q) | Q(email__icontains=q)
        ).order_by("name")[:5]
        customers = [
            {"id": c.id, "name": c.name, "phone": c.phone, "url": reverse("customers:detail", args=[c.id])}
            for c in customer_qs
        ]
    return JsonResponse({"products": products, "customers": customers})


# --------------------------------------------------------------------------
# Notifications
# --------------------------------------------------------------------------

@login_required
def notifications(request):
    notes = (
        Notification.objects.filter(recipient=request.user)
        .select_related("branch")[:100]
    )
    return render(request, "core/notifications.html", {
        "active_nav": "notifications",
        "notifications": notes,
    })


@login_required
@require_POST
def notification_read(request, pk):
    note = get_object_or_404(Notification, pk=pk, recipient=request.user)
    note.mark_read()
    if note.url:
        return redirect(note.url)
    return redirect("core:notifications")


@login_required
@require_POST
def notifications_read_all(request):
    updated = Notification.objects.filter(recipient=request.user, is_read=False).update(is_read=True)
    messages.success(request, f"Marked {updated} notification(s) as read.")
    return redirect("core:notifications")


# --------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------

def _report_range(request):
    """Resolve the ?from / ?to / ?branch query into (start, end, branch)."""
    today = timezone.localdate()

    def parse(value, fallback):
        if not value:
            return fallback
        try:
            return timezone.datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return fallback

    start = parse(request.GET.get("from"), today - timedelta(days=29))
    end = parse(request.GET.get("to"), today)
    if start > end:
        start, end = end, start

    branch = None
    branch_id = request.GET.get("branch")
    if branch_id:
        branch = Branch.objects.filter(id=branch_id).first()
    elif not request.GET:
        # Default to the viewer's own branch on first load.
        profile = getattr(request.user, "profile", None)
        branch = profile.branch if profile else None
    return start, end, branch


def _report_data(request):
    start, end, branch = _report_range(request)

    sales = Sale.objects.filter(
        status=Sale.STATUS_COMPLETED, completed_at__date__gte=start, completed_at__date__lte=end
    )
    voids = Sale.objects.filter(
        status=Sale.STATUS_VOID, voided_at__date__gte=start, voided_at__date__lte=end
    )
    if branch:
        sales = sales.filter(branch=branch)
        voids = voids.filter(branch=branch)

    totals = sales.aggregate(
        revenue=Coalesce(Sum("total"), Decimal("0"), output_field=MONEY),
        net_sales=Coalesce(Sum("subtotal"), Decimal("0"), output_field=MONEY),
        tax=Coalesce(Sum("tax"), Decimal("0"), output_field=MONEY),
        order_discount=Coalesce(Sum("discount_amount"), Decimal("0"), output_field=MONEY),
        count=Count("id"),
    )
    revenue, net_sales, sale_count = totals["revenue"], totals["net_sales"], totals["count"]

    line_items = SaleItem.objects.filter(sale__in=sales)
    item_agg = line_items.aggregate(
        units=Coalesce(Sum("quantity"), Decimal("0"), output_field=MONEY),
        line_discount=Coalesce(Sum("discount_amount"), Decimal("0"), output_field=MONEY),
        cost=Coalesce(Sum(F("quantity") * F("unit_cost"), output_field=MONEY), Decimal("0"), output_field=MONEY),
    )
    cost_of_goods = item_agg["cost"]
    gross_profit = net_sales - cost_of_goods
    total_discount = totals["order_discount"] + item_agg["line_discount"]

    by_day_map = {
        row["day"]: row
        for row in sales.annotate(day=TruncDate("completed_at"))
        .values("day")
        .annotate(revenue=Coalesce(Sum("total"), Decimal("0"), output_field=MONEY), count=Count("id"))
    }
    peak = max([row["revenue"] for row in by_day_map.values()] or [Decimal("0")])
    by_day = []
    day = start
    while day <= end:
        row = by_day_map.get(day, {"revenue": Decimal("0"), "count": 0})
        percent = int((row["revenue"] / peak) * 100) if peak else 0
        by_day.append({"day": day, "revenue": row["revenue"], "count": row["count"], "bar_percent": percent})
        day += timedelta(days=1)

    # Per-product performance for the period: revenue, cost, profit and margin,
    # so the same query set can drive both the "best sellers" and "losing money"
    # views instead of two separate passes over the sales.
    product_rows = list(
        line_items.filter(product__isnull=False)
        .values("product__id", "product__sku", "product__name")
        .annotate(
            units=Coalesce(Sum("quantity"), Decimal("0"), output_field=MONEY),
            revenue=Coalesce(Sum("line_total"), Decimal("0"), output_field=MONEY),
            cost=Coalesce(Sum(F("quantity") * F("unit_cost"), output_field=MONEY), Decimal("0"), output_field=MONEY),
            discount=Coalesce(Sum("discount_amount"), Decimal("0"), output_field=MONEY),
        )
    )
    for row in product_rows:
        row["profit"] = row["revenue"] - row["cost"]
        row["margin_percent"] = (row["profit"] / row["revenue"] * 100) if row["revenue"] else None

    top_products = sorted(product_rows, key=lambda r: r["revenue"], reverse=True)[:8]
    # "Losing money" = actually unprofitable this period (needs a real cost on
    # file — rows with cost=0 just mean cost was never entered, not that the
    # item is free to sell).
    losing_products = sorted(
        (r for r in product_rows if r["cost"] > 0 and r["profit"] < 0),
        key=lambda r: r["profit"],
    )[:8]
    for row in losing_products:
        row["loss"] = -row["profit"]
    thin_margin_products = sorted(
        (r for r in product_rows if r["cost"] > 0 and r["profit"] >= 0),
        key=lambda r: (r["margin_percent"] if r["margin_percent"] is not None else 999),
    )[:8]

    payment_labels = dict(Payment.METHOD_CHOICES)
    payment_mix = [
        {"label": payment_labels.get(row["method"], row["method"]), "count": row["count"], "amount": row["amount"]}
        for row in Payment.objects.filter(sale__in=sales)
        .values("method")
        .annotate(count=Count("id"), amount=Coalesce(Sum("amount"), Decimal("0"), output_field=MONEY))
        .order_by("-amount")
    ]

    by_cashier = list(
        sales.values("cashier__username")
        .annotate(count=Count("id"), revenue=Coalesce(Sum("total"), Decimal("0"), output_field=MONEY))
        .order_by("-revenue")
    )
    by_branch = list(
        sales.values("branch__name")
        .annotate(count=Count("id"), revenue=Coalesce(Sum("total"), Decimal("0"), output_field=MONEY))
        .order_by("-revenue")
    )

    stock_qs = StockLevel.objects.filter(product__is_active=True)
    if branch:
        stock_qs = stock_qs.filter(branch=branch)

    void_agg = voids.aggregate(
        count=Count("id"), value=Coalesce(Sum("total"), Decimal("0"), output_field=MONEY)
    )

    dead_stock, dead_stock_days = _dead_stock_rows(branch)

    return {
        "active_nav": "reports",
        "start": start,
        "end": end,
        "day_span": (end - start).days + 1,
        "branch": branch,
        "branches": Branch.objects.filter(is_active=True, is_warehouse=False),
        "revenue": revenue,
        "net_sales": net_sales,
        "tax": totals["tax"],
        "sale_count": sale_count,
        "average_sale": (revenue / sale_count) if sale_count else Decimal("0"),
        "units_sold": item_agg["units"],
        "tint_jobs": sales.filter(items__color_formula__isnull=False).distinct().count(),
        "void_count": void_agg["count"],
        "void_value": void_agg["value"],
        # Profit & loss
        "cost_of_goods": cost_of_goods,
        "gross_profit": gross_profit,
        "gross_margin_percent": (gross_profit / net_sales * 100) if net_sales else None,
        "order_discount_total": totals["order_discount"],
        "line_discount_total": item_agg["line_discount"],
        "total_discount": total_discount,
        "is_loss_period": gross_profit < 0,
        "by_day": by_day,
        "top_products": top_products,
        "losing_products": losing_products,
        "thin_margin_products": thin_margin_products,
        "has_cost_data": any(r["cost"] > 0 for r in product_rows),
        "payment_mix": payment_mix,
        "by_cashier": by_cashier,
        "by_branch": by_branch,
        "low_stock_count": stock_qs.filter(quantity__lte=F("reorder_threshold"), quantity__gt=0).count(),
        "out_of_stock_count": stock_qs.filter(quantity=0).count(),
        "dead_stock": dead_stock,
        "dead_stock_days": dead_stock_days,
        "dead_stock_value": sum((row["tied_up_value"] for row in dead_stock), start=Decimal("0")),
    }


def _dead_stock_rows(branch):
    """Products sitting on the shelf that haven't sold in a while (or ever).

    Scoped to whatever stock/sale data exists regardless of the report's date
    range — "how long since this last sold" needs the real last-sale date,
    not just whether it sold within the selected window.
    """
    from django.conf import settings

    dead_days = int(getattr(settings, "DEAD_STOCK_DAYS", 30))
    today = timezone.localdate()

    stock_qs = StockLevel.objects.filter(product__is_active=True, quantity__gt=0)
    if branch:
        stock_qs = stock_qs.filter(branch=branch)
    stock_by_product = {}
    for row in stock_qs.values("product_id").annotate(qty=Sum("quantity")):
        stock_by_product[row["product_id"]] = row["qty"] or 0
    if not stock_by_product:
        return [], dead_days

    sale_item_qs = SaleItem.objects.filter(sale__status=Sale.STATUS_COMPLETED, product__isnull=False)
    if branch:
        sale_item_qs = sale_item_qs.filter(sale__branch=branch)
    last_sold_map = {
        row["product_id"]: row["last"]
        for row in sale_item_qs.values("product_id").annotate(last=Max("sale__completed_at"))
    }

    products = Product.objects.filter(id__in=stock_by_product.keys()).select_related("category", "brand")
    rows = []
    for product in products:
        stock_qty = stock_by_product.get(product.id, 0)
        last_sold_at = last_sold_map.get(product.id)
        never_sold = last_sold_at is None
        if never_sold:
            days_idle = (today - product.created_at.date()).days
            last_sold_date = None
        else:
            last_sold_date = timezone.localtime(last_sold_at).date()
            days_idle = (today - last_sold_date).days
        if not never_sold and days_idle < dead_days:
            continue
        unit_value = product.cost_price or product.unit_price
        rows.append({
            "product": product,
            "stock_qty": stock_qty,
            "last_sold": last_sold_date,
            "days_idle": days_idle,
            "never_sold": never_sold,
            "tied_up_value": (Decimal(stock_qty) * unit_value).quantize(Decimal("0.01")),
        })
    rows.sort(key=lambda r: (r["never_sold"], r["days_idle"]), reverse=True)
    return rows, dead_days


@login_required
@manager_required
def reports(request):
    return render(request, "core/reports.html", _report_data(request))


def _export_scope_and_filename(data, ext):
    import re

    scope_raw = data["branch"].code if data["branch"] else "ALL"
    # Branch codes are free text — strip anything that isn't safe in a
    # filename (quotes, slashes, etc. can mangle the Content-Disposition
    # header or fail to save on some operating systems).
    scope = re.sub(r"[^A-Za-z0-9_-]+", "_", scope_raw).strip("_") or "ALL"
    filename = f"chromapos-{scope}-{data['start']:%Y%m%d}-{data['end']:%Y%m%d}.{ext}"
    return scope, filename


def _export_csv(data):
    _, filename = _export_scope_and_filename(data, "csv")
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)

    writer.writerow(["ChromaPOS sales report"])
    writer.writerow(["Branch", data["branch"].name if data["branch"] else "All branches"])
    writer.writerow(["From", data["start"], "To", data["end"]])
    writer.writerow([])
    writer.writerow(["Revenue", data["revenue"], "Tax", data["tax"], "Sales", data["sale_count"]])
    writer.writerow(["Units sold", data["units_sold"], "Tint jobs", data["tint_jobs"]])
    writer.writerow(["Voided sales", data["void_count"], "Voided value", data["void_value"]])
    writer.writerow([])

    writer.writerow(["Profit & Loss", ""])
    writer.writerow(["Net sales (after discounts, before tax)", data["net_sales"]])
    writer.writerow(["Cost of goods sold", data["cost_of_goods"]])
    writer.writerow(["Gross profit", data["gross_profit"]])
    writer.writerow(["Gross margin %", f"{data['gross_margin_percent']:.1f}" if data["gross_margin_percent"] is not None else "n/a"])
    writer.writerow(["Total discounts given", data["total_discount"]])
    writer.writerow(["  Order-level discounts", data["order_discount_total"]])
    writer.writerow(["  Line-item discounts", data["line_discount_total"]])
    if not data["has_cost_data"]:
        writer.writerow(["Note", "No product cost prices on file yet — profit figures are incomplete."])
    writer.writerow([])

    writer.writerow(["Date", "Sales", "Revenue"])
    for row in data["by_day"]:
        writer.writerow([row["day"], row["count"], row["revenue"]])
    writer.writerow([])

    writer.writerow(["Best-selling products", ""])
    writer.writerow(["SKU", "Product", "Units", "Revenue", "Cost", "Profit", "Margin %"])
    for row in data["top_products"]:
        writer.writerow([
            row["product__sku"], row["product__name"], row["units"], row["revenue"],
            row["cost"], row["profit"],
            f"{row['margin_percent']:.1f}" if row["margin_percent"] is not None else "n/a",
        ])
    writer.writerow([])

    if data["losing_products"]:
        writer.writerow(["Products losing money this period", ""])
        writer.writerow(["SKU", "Product", "Units", "Revenue", "Cost", "Loss"])
        for row in data["losing_products"]:
            writer.writerow([row["product__sku"], row["product__name"], row["units"], row["revenue"], row["cost"], row["profit"]])
        writer.writerow([])

    if data["dead_stock"]:
        writer.writerow([f"Slow-moving stock (no sale in {data['dead_stock_days']}+ days)", ""])
        writer.writerow(["SKU", "Product", "In Stock", "Last Sold", "Days Since Sale", "Value Tied Up"])
        for row in data["dead_stock"]:
            writer.writerow([
                row["product"].sku, row["product"].name, row["stock_qty"],
                row["last_sold"] if row["last_sold"] else "Never sold",
                row["days_idle"], row["tied_up_value"],
            ])
        writer.writerow([])

    writer.writerow(["Payment method", "Count", "Amount"])
    for row in data["payment_mix"]:
        writer.writerow([row["label"], row["count"], row["amount"]])
    return response


def _export_xlsx(data):
    _, filename = _export_scope_and_filename(data, "xlsx")
    branch_label = data["branch"].name if data["branch"] else "All branches"

    wb = Workbook()
    bold = Font(bold=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill_kwargs = {"fill_type": "solid", "fgColor": "004AC6"}

    def style_header_row(ws, row_idx, col_count):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = header_font
            cell.fill = PatternFill(**header_fill_kwargs)

    def autosize(ws):
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 40)

    summary = wb.active
    summary.title = "Summary"
    summary.append(["ChromaPOS Sales Report"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Branch", branch_label])
    summary.append(["From", str(data["start"]), "To", str(data["end"])])
    summary.append([])
    summary.append(["Metric", "Value"])
    style_header_row(summary, summary.max_row, 2)
    for label, value in [
        ("Revenue", float(data["revenue"])),
        ("Tax", float(data["tax"])),
        ("Sales", data["sale_count"]),
        ("Average Sale", float(data["average_sale"])),
        ("Units Sold", float(data["units_sold"])),
        ("Tint Jobs", data["tint_jobs"]),
        ("Voided Sales", data["void_count"]),
        ("Voided Value", float(data["void_value"])),
        ("Net Sales (after discounts)", float(data["net_sales"])),
        ("Cost of Goods Sold", float(data["cost_of_goods"])),
        ("Gross Profit", float(data["gross_profit"])),
        ("Gross Margin %", round(float(data["gross_margin_percent"]), 1) if data["gross_margin_percent"] is not None else "n/a"),
        ("Total Discounts Given", float(data["total_discount"])),
        ("  Order-level Discounts", float(data["order_discount_total"])),
        ("  Line-item Discounts", float(data["line_discount_total"])),
    ]:
        summary.append([label, value])
    autosize(summary)

    by_day = wb.create_sheet("Revenue by Day")
    by_day.append(["Date", "Sales", "Revenue"])
    style_header_row(by_day, 1, 3)
    for row in data["by_day"]:
        by_day.append([str(row["day"]), row["count"], float(row["revenue"])])
    autosize(by_day)

    top_products = wb.create_sheet("Top Products")
    top_products.append(["SKU", "Product", "Units", "Revenue", "Cost", "Profit", "Margin %"])
    style_header_row(top_products, 1, 7)
    for row in data["top_products"]:
        top_products.append([
            row["product__sku"], row["product__name"], float(row["units"]), float(row["revenue"]),
            float(row["cost"]), float(row["profit"]),
            round(float(row["margin_percent"]), 1) if row["margin_percent"] is not None else "n/a",
        ])
    autosize(top_products)

    if data["losing_products"]:
        losing = wb.create_sheet("Losing Money")
        losing.append(["SKU", "Product", "Units", "Revenue", "Cost", "Loss"])
        style_header_row(losing, 1, 6)
        for row in data["losing_products"]:
            losing.append([row["product__sku"], row["product__name"], float(row["units"]), float(row["revenue"]), float(row["cost"]), float(row["profit"])])
        autosize(losing)

    if data["dead_stock"]:
        dead = wb.create_sheet("Slow-Moving Stock")
        dead.append(["SKU", "Product", "In Stock", "Last Sold", "Days Since Sale", "Value Tied Up"])
        style_header_row(dead, 1, 6)
        for row in data["dead_stock"]:
            dead.append([
                row["product"].sku, row["product"].name, row["stock_qty"],
                str(row["last_sold"]) if row["last_sold"] else "Never sold",
                row["days_idle"], float(row["tied_up_value"]),
            ])
        autosize(dead)

    payments = wb.create_sheet("Payment Mix")
    payments.append(["Payment Method", "Count", "Amount"])
    style_header_row(payments, 1, 3)
    for row in data["payment_mix"]:
        payments.append([row["label"], row["count"], float(row["amount"])])
    autosize(payments)

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _export_pdf(data):
    from xml.sax.saxutils import escape as _xml_escape

    _, filename = _export_scope_and_filename(data, "pdf")
    branch_label = _xml_escape(data["branch"].name if data["branch"] else "All branches")

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("ChromaPOS Sales Report", styles["Title"]),
        Paragraph(
            f"{branch_label} · {data['start']:%b %d, %Y} – {data['end']:%b %d, %Y}",
            styles["Normal"],
        ),
        Spacer(1, 0.25 * inch),
    ]

    header_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004AC6")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c3c6d7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f5")]),
    ])

    summary_rows = [
        ["Revenue", f"${data['revenue']:.2f}", "Tax", f"${data['tax']:.2f}"],
        ["Sales", str(data["sale_count"]), "Average Sale", f"${data['average_sale']:.2f}"],
        ["Units Sold", f"{data['units_sold']:.0f}", "Tint Jobs", str(data["tint_jobs"])],
        ["Voided Sales", str(data["void_count"]), "Voided Value", f"${data['void_value']:.2f}"],
    ]
    summary_table = Table(summary_rows, colWidths=[1.6 * inch, 1.7 * inch, 1.6 * inch, 1.7 * inch])
    summary_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c3c6d7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [Paragraph("Summary", styles["Heading2"]), summary_table, Spacer(1, 0.25 * inch)]

    margin_text = f"{data['gross_margin_percent']:.1f}%" if data["gross_margin_percent"] is not None else "n/a"
    pl_rows = [
        ["Net Sales", f"${data['net_sales']:.2f}", "Cost of Goods", f"${data['cost_of_goods']:.2f}"],
        ["Gross Profit", f"${data['gross_profit']:.2f}", "Gross Margin", margin_text],
        ["Total Discounts", f"${data['total_discount']:.2f}", "  (order + line item)", ""],
    ]
    pl_table = Table(pl_rows, colWidths=[1.6 * inch, 1.7 * inch, 1.6 * inch, 1.7 * inch])
    pl_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c3c6d7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#c0392b") if data["is_loss_period"] else colors.HexColor("#1a7f37")),
    ]))
    story += [Paragraph("Profit &amp; Loss", styles["Heading2"]), pl_table]
    if not data["has_cost_data"]:
        story += [Spacer(1, 0.08 * inch), Paragraph(
            "<i>No product cost prices are on file yet, so profit figures above are incomplete.</i>",
            styles["Normal"],
        )]
    story += [Spacer(1, 0.25 * inch)]

    if data["top_products"]:
        rows = [["SKU", "Product", "Units", "Revenue", "Profit", "Margin"]] + [
            [
                row["product__sku"], row["product__name"], f"{row['units']:.0f}", f"${row['revenue']:.2f}",
                f"${row['profit']:.2f}", f"{row['margin_percent']:.1f}%" if row["margin_percent"] is not None else "n/a",
            ]
            for row in data["top_products"]
        ]
        table = Table(rows, colWidths=[1 * inch, 2 * inch, 0.7 * inch, 1 * inch, 0.9 * inch, 0.8 * inch], repeatRows=1)
        table.setStyle(header_style)
        story += [Paragraph("Best-Selling Products", styles["Heading2"]), table, Spacer(1, 0.25 * inch)]

    if data["losing_products"]:
        rows = [["SKU", "Product", "Units", "Revenue", "Cost", "Loss"]] + [
            [row["product__sku"], row["product__name"], f"{row['units']:.0f}", f"${row['revenue']:.2f}", f"${row['cost']:.2f}", f"-${abs(row['profit']):.2f}"]
            for row in data["losing_products"]
        ]
        table = Table(rows, colWidths=[1 * inch, 2.2 * inch, 0.7 * inch, 1 * inch, 0.9 * inch, 0.9 * inch], repeatRows=1)
        loss_style = TableStyle(header_style.getCommands())
        loss_style.add("TEXTCOLOR", (5, 1), (5, -1), colors.HexColor("#c0392b"))
        table.setStyle(loss_style)
        story += [Paragraph("Products Losing Money", styles["Heading2"]), table, Spacer(1, 0.25 * inch)]

    if data["dead_stock"]:
        # PDF table layout is the slowest of the three export formats — cap it
        # to the biggest-value rows so a large catalog can't blow up
        # generation time (or a request timeout somewhere in front of the
        # app) into a download that never finishes writing. CSV/Excel always
        # carry the complete list.
        PDF_DEAD_STOCK_CAP = 60
        dead_stock_rows = sorted(data["dead_stock"], key=lambda r: r["tied_up_value"], reverse=True)
        truncated = len(dead_stock_rows) > PDF_DEAD_STOCK_CAP
        dead_stock_rows = dead_stock_rows[:PDF_DEAD_STOCK_CAP]
        rows = [["SKU", "Product", "In Stock", "Last Sold", "Days Idle", "Value Tied Up"]] + [
            [
                row["product"].sku, row["product"].name, str(row["stock_qty"]),
                row["last_sold"].strftime("%b %d, %Y") if row["last_sold"] else "Never sold",
                str(row["days_idle"]), f"${row['tied_up_value']:.2f}",
            ]
            for row in dead_stock_rows
        ]
        table = Table(rows, colWidths=[1 * inch, 1.9 * inch, 0.7 * inch, 1 * inch, 0.7 * inch, 1 * inch], repeatRows=1)
        table.setStyle(header_style)
        heading = (
            f"Slow-Moving Stock (no sale in {data['dead_stock_days']}+ days) "
            f"— {len(data['dead_stock'])} product(s), ${data['dead_stock_value']:.2f} tied up"
        )
        story += [Paragraph(heading, styles["Heading2"]), table]
        if truncated:
            story += [Spacer(1, 0.06 * inch), Paragraph(
                f"<i>Showing the {PDF_DEAD_STOCK_CAP} highest-value rows. "
                f"Export as CSV or Excel for the complete list of {len(data['dead_stock'])}.</i>",
                styles["Normal"],
            )]
        story += [Spacer(1, 0.25 * inch)]

    if data["payment_mix"]:
        rows = [["Payment Method", "Count", "Amount"]] + [
            [row["label"], str(row["count"]), f"${row['amount']:.2f}"] for row in data["payment_mix"]
        ]
        table = Table(rows, colWidths=[2.5 * inch, 1.5 * inch, 2.4 * inch], repeatRows=1)
        table.setStyle(header_style)
        story += [Paragraph("Payment Mix", styles["Heading2"]), table, Spacer(1, 0.25 * inch)]

    if data["by_day"]:
        # Same reasoning as Slow-Moving Stock above: a very wide date range
        # (a year-plus) makes for a huge table. Cap it to the most recent
        # stretch and point to CSV/Excel for the full daily series.
        PDF_BY_DAY_CAP = 62
        by_day_rows = data["by_day"]
        truncated = len(by_day_rows) > PDF_BY_DAY_CAP
        if truncated:
            by_day_rows = by_day_rows[-PDF_BY_DAY_CAP:]
        rows = [["Date", "Sales", "Revenue"]] + [
            [str(row["day"]), str(row["count"]), f"${row['revenue']:.2f}"] for row in by_day_rows
        ]
        table = Table(rows, colWidths=[2.5 * inch, 1.5 * inch, 2.4 * inch], repeatRows=1)
        table.setStyle(header_style)
        heading = "Revenue by Day"
        if truncated:
            heading += f" (most recent {PDF_BY_DAY_CAP} of {len(data['by_day'])} days)"
        story += [Paragraph(heading, styles["Heading2"]), table]
        if truncated:
            story += [Spacer(1, 0.06 * inch), Paragraph(
                "<i>Export as CSV or Excel for the complete daily series.</i>", styles["Normal"],
            )]

    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
@manager_required
def reports_export(request):
    import logging

    data = _report_data(request)
    fmt = (request.GET.get("format") or "csv").lower()
    try:
        if fmt == "xlsx":
            return _export_xlsx(data)
        if fmt == "pdf":
            return _export_pdf(data)
        return _export_csv(data)
    except Exception:
        # Never hand back a half-built file with the right extension but
        # broken insides — that's worse than an error, since it looks like
        # it worked until the person tries to open it. Fall back to CSV
        # (plain text, always openable) and log the real cause for us.
        logging.getLogger(__name__).exception(
            "Report export failed (format=%s, branch=%s, %s to %s) — falling back to CSV.",
            fmt, data.get("branch"), data.get("start"), data.get("end"),
        )
        response = _export_csv(data)
        response["X-ChromaPOS-Export-Fallback"] = "1"
        return response